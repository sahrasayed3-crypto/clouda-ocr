from __future__ import annotations

import io
import json
import subprocess
import sys
from pathlib import Path

import fitz
import pytest
from docx import Document
from PIL import Image
from pypdf.errors import PdfReadError

from pdfword.docx_export import markdown_to_docx
from pdfword.engines import (
    DIRECT_TEXT_ENGINE,
    DirectPdfTextEngine,
    OCRResult,
    OCR_STATUS_FAILED,
    OCR_STATUS_PENDING_MODEL,
    OCR_STATUS_SUCCEEDED,
)
from pdfword.ocr_pipeline import BLANK_PAGE_ROUTE, NEAR_BLANK_PAGE_ROUTE, process_pdf

FIXTURES = Path(__file__).with_name("fixtures")
ROOT = Path(__file__).resolve().parents[1]


def _process(name: str):
    return process_pdf(
        (FIXTURES / name).read_bytes(),
        from_page=None,
        to_page=None,
        page_numbers=[1],
        progress_bar=None,
        status_placeholder=None,
    )


def test_copyright_free_pdf_fixtures_are_present() -> None:
    expected = {
        "digital_text.pdf",
        "scanned.pdf",
        "blank.pdf",
        "near_blank_page_number.pdf",
        "near_blank_stamp.pdf",
        "mixed.pdf",
        "corrupt.pdf",
        "empty.pdf",
        "not_a_pdf.txt",
    }
    assert expected <= {path.name for path in FIXTURES.iterdir()}


def test_digital_pdf_extracts_text_without_ocr() -> None:
    rows, text = _process("digital_text.pdf")

    assert text.startswith("Digital PDF text")
    assert rows[0].route_used == "direct_pdf_text"
    assert rows[0].metadata["page_state"] == "digital_text"
    assert rows[0].metadata["embedded_text_chars"] > 20
    assert rows[0].metadata["document_intelligence"]["decision"] == (
        "trusted_digital_text"
    )
    assert rows[0].metadata["document_intelligence"]["gate_verdict"] == "trusted"


def _hidden_partial_text_pdf() -> bytes:
    image = Image.new("RGB", (1190, 1684), "white")
    image_output = io.BytesIO()
    image.save(image_output, format="PNG")
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=image_output.getvalue())
    page.insert_text((72, 72), "Header only", fontsize=6, fontname="helv")
    payload = document.tobytes(garbage=4, deflate=True)
    document.close()
    return payload


def test_direct_engine_does_not_fabricate_confidence() -> None:
    result = DirectPdfTextEngine().extract_page(
        pdf_bytes=(FIXTURES / "digital_text.pdf").read_bytes(), page_no=1
    )

    assert result.success
    assert result.confidence is None


def test_direct_text_digest_mismatch_becomes_visible_review(monkeypatch) -> None:
    monkeypatch.setattr(
        DIRECT_TEXT_ENGINE,
        "extract_page",
        lambda **_kwargs: OCRResult(
            engine_name="direct_pdf_text",
            status=OCR_STATUS_SUCCEEDED,
            text="changed after analysis",
            confidence=None,
        ),
    )

    rows, text = _process("digital_text.pdf")

    assert rows[0].route_used == "review_required"
    assert rows[0].requires_manual_review is True
    assert "changed after analysis" not in rows[0].markdown
    assert "changed after analysis" not in text
    assert "REQUIRES REVIEW" in rows[0].markdown


def test_review_placeholder_is_visible_in_docx_without_quality_percentage(
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        DIRECT_TEXT_ENGINE,
        "extract_page",
        lambda **_kwargs: OCRResult(
            engine_name="direct_pdf_text",
            status=OCR_STATUS_SUCCEEDED,
            text="mismatched private text",
            confidence=None,
        ),
    )
    rows, _ = _process("digital_text.pdf")

    document = Document(io.BytesIO(markdown_to_docx(rows)))
    visible = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "REQUIRES REVIEW" in visible
    assert "mismatched private text" not in visible
    assert "%" not in visible
    assert "estimated quality" not in visible.lower()


def test_review_diagnostics_are_categorical_and_do_not_expose_page_text(
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        DIRECT_TEXT_ENGINE,
        "extract_page",
        lambda **_kwargs: OCRResult(
            engine_name="direct_pdf_text",
            status=OCR_STATUS_SUCCEEDED,
            text="mismatch",
            confidence=None,
        ),
    )
    rows, _ = _process("digital_text.pdf")

    diagnostics = rows[0].metadata["document_intelligence"]
    serialized = json.dumps(diagnostics).lower()

    assert diagnostics["decision"] == "review_required"
    assert diagnostics["review_required"] is True
    assert "digital pdf text" not in serialized
    assert "accuracy" not in serialized
    assert "confidence" not in serialized
    assert "quality" not in serialized


def test_untrusted_hidden_text_is_not_emitted_as_conversion_output() -> None:
    rows, text = process_pdf(
        _hidden_partial_text_pdf(),
        from_page=1,
        to_page=1,
        progress_bar=None,
        status_placeholder=None,
    )

    assert rows[0].route_used == OCR_STATUS_PENDING_MODEL
    metadata = rows[0].metadata
    assert metadata is not None
    diagnostics = metadata["document_intelligence"]
    assert isinstance(diagnostics, dict)
    assert diagnostics["decision"] == "ocr_required"
    assert "Header only" not in rows[0].markdown
    assert "Header only" not in text


@pytest.mark.parametrize("name", ["scanned.pdf"])
def test_scanned_pdf_waits_for_an_approved_ocr_model(name: str) -> None:
    rows, _ = _process(name)

    assert rows[0].route_used == OCR_STATUS_PENDING_MODEL
    assert rows[0].accepted is False
    assert rows[0].requires_manual_review is True
    assert rows[0].metadata["page_state"] == OCR_STATUS_PENDING_MODEL
    assert rows[0].metadata["embedded_image_count"] == 1


def test_blank_page_is_not_mislabelled_as_needing_ocr() -> None:
    rows, text = _process("blank.pdf")

    assert text == ""
    assert rows[0].route_used == BLANK_PAGE_ROUTE
    assert rows[0].metadata["page_state"] == BLANK_PAGE_ROUTE
    assert rows[0].requires_manual_review is False


@pytest.mark.parametrize("name", ["near_blank_page_number.pdf", "near_blank_stamp.pdf"])
def test_near_blank_content_is_preserved_for_review(name: str) -> None:
    rows, _ = _process(name)

    assert rows[0].route_used == NEAR_BLANK_PAGE_ROUTE
    assert rows[0].requires_manual_review is True
    assert rows[0].metadata["page_state"] == NEAR_BLANK_PAGE_ROUTE


def test_page_number_text_is_not_deleted_when_near_blank() -> None:
    rows, _ = _process("near_blank_page_number.pdf")

    assert rows[0].markdown == "12"
    document = Document(io.BytesIO(markdown_to_docx(rows)))
    assert "12" in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_mixed_pdf_preserves_page_order_and_status_metadata() -> None:
    rows, _ = process_pdf(
        (FIXTURES / "mixed.pdf").read_bytes(),
        from_page=1,
        to_page=2,
        progress_bar=None,
        status_placeholder=None,
    )

    assert [row.page_no for row in rows] == [1, 2]
    assert [row.route_used for row in rows] == [
        "direct_pdf_text",
        OCR_STATUS_PENDING_MODEL,
    ]
    assert all(row.metadata and row.metadata["page_state"] for row in rows)


@pytest.mark.parametrize("name", ["empty.pdf", "corrupt.pdf"])
def test_empty_or_corrupt_pdf_fails_validation(name: str) -> None:
    with pytest.raises(PdfReadError):
        _process(name)


def test_non_pdf_is_rejected_by_the_direct_engine_without_a_crash() -> None:
    result = DirectPdfTextEngine().extract_page(
        pdf_bytes=(FIXTURES / "not_a_pdf.txt").read_bytes(), page_no=1
    )

    assert result.status == OCR_STATUS_FAILED
    assert result.failure_reason


def test_pending_model_is_neither_success_nor_final_failure() -> None:
    rows, _ = _process("scanned.pdf")

    assert rows[0].route_used == OCR_STATUS_PENDING_MODEL
    assert rows[0].accepted is False
    assert rows[0].requires_manual_review is True


def test_demo_creates_docx_and_page_metadata(tmp_path: Path) -> None:
    result = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "demo.py"),
            "--output-dir",
            str(tmp_path),
        ],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    )

    payload = json.loads((tmp_path / "page_statuses.json").read_text(encoding="utf-8"))
    assert (tmp_path / "digital_text.docx").is_file()
    assert [page["route_used"] for page in payload["pages"]] == [
        "direct_pdf_text",
        OCR_STATUS_PENDING_MODEL,
        BLANK_PAGE_ROUTE,
    ]
    assert "States:" in result.stdout
