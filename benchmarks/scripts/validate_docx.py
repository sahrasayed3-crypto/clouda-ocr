from __future__ import annotations

import io
import zipfile
from pathlib import Path

from docx import Document


def validate_docx_bytes(payload: bytes, *, expected_pages: int = 1) -> dict:
    errors: list[str] = []
    paragraphs: list[str] = []
    has_bidi = False
    has_page_break = False
    media_count = 0
    table_count = 0

    try:
        document = Document(io.BytesIO(payload))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_count = len(document.tables)
    except Exception as exc:
        return {
            "valid": False,
            "errors": [
                f"python-docx failed to open output: {type(exc).__name__}: {exc}"
            ],
            "paragraph_count": 0,
            "table_count": 0,
            "media_count": 0,
            "has_page_break": False,
            "has_rtl": False,
        }

    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            names = archive.namelist()
            media_count = len(
                [name for name in names if name.startswith("word/media/")]
            )
            document_xml = archive.read("word/document.xml").decode(
                "utf-8", errors="replace"
            )
            has_bidi = "w:bidi" in document_xml
            has_page_break = 'w:type="page"' in document_xml
            if "<w:tbl>" in document_xml:
                table_count = max(1, table_count)
    except Exception as exc:
        errors.append(f"DOCX zip/XML inspection failed: {type(exc).__name__}: {exc}")

    if expected_pages > 1 and not has_page_break:
        errors.append("Expected page breaks were not found.")
    if table_count:
        errors.append("Text-only DOCX contains tables.")
    if media_count:
        errors.append("Text-only DOCX contains embedded media.")

    return {
        "valid": not errors,
        "errors": errors,
        "paragraph_count": len(paragraphs),
        "non_empty_paragraph_count": len([text for text in paragraphs if text.strip()]),
        "table_count": table_count,
        "media_count": media_count,
        "has_page_break": has_page_break,
        "has_rtl": has_bidi,
        "first_paragraph": next((text for text in paragraphs if text.strip()), ""),
    }


def validate_docx_path(path: str | Path, *, expected_pages: int = 1) -> dict:
    return validate_docx_bytes(Path(path).read_bytes(), expected_pages=expected_pages)
