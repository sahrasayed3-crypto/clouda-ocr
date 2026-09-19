"""Tiny offline end-to-end release check for the canonical PDF pipeline."""

from __future__ import annotations

import io
from dataclasses import dataclass

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
)

from .docx_export import markdown_to_docx
from .ocr_pipeline import process_pdf


@dataclass(frozen=True)
class ReleaseSelfTestResult:
    ok: bool
    states: tuple[str, ...]
    page_boundaries_preserved: bool
    user_facing_percentages: bool


def _stream(data: bytes) -> DecodedStreamObject:
    value = DecodedStreamObject()
    value.set_data(data)
    return value


def _fixture_pdf() -> bytes:
    writer = PdfWriter()
    text = (
        "BT /F1 12 Tf 36 220 Td "
        "(Trusted digital text for the deterministic offline release self test. "
        "It is deliberately long enough to provide structural evidence.) Tj ET"
    ).encode("ascii")
    first = writer.add_blank_page(width=300, height=300)
    first[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): DictionaryObject(
                        {
                            NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica"),
                        }
                    )
                }
            )
        }
    )
    first[NameObject("/Contents")] = writer._add_object(_stream(text))
    writer.add_blank_page(width=300, height=300)
    image_page = writer.add_blank_page(width=300, height=300)
    # A page-sized image is structural evidence of a scan; a one-pixel image
    # is intentionally classified near-blank by the canonical analyzer.
    image = _stream(bytes((255,)) * (300 * 300 * 3))
    image.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(300),
            NameObject("/Height"): NumberObject(300),
            NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
            NameObject("/BitsPerComponent"): NumberObject(8),
        }
    )
    image_page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/XObject"): DictionaryObject(
                {NameObject("/Im1"): writer._add_object(image)}
            )
        }
    )
    image_page[NameObject("/Contents")] = writer._add_object(
        _stream(b"q 300 0 0 300 0 0 cm /Im1 Do Q")
    )
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def run_release_self_test() -> ReleaseSelfTestResult:
    """Exercise canonical trusted, blank, and OCR-required outcomes offline."""
    pages, _ = process_pdf(
        _fixture_pdf(),
        from_page=1,
        to_page=3,
        progress_bar=None,
        status_placeholder=None,
    )
    docx = Document(io.BytesIO(markdown_to_docx(pages)))
    visible = "\n".join(paragraph.text for paragraph in docx.paragraphs)
    states = tuple(
        str((page.metadata or {}).get("page_state", page.route_used)) for page in pages
    )
    percentages = any("%" in paragraph.text for paragraph in docx.paragraphs)
    boundaries = (
        len(pages) == 3
        and len(docx.paragraphs) >= 3
        and "Requires future OCR model" in visible
    )
    required_states = {"digital_text", "blank_page", "pending_ocr_model"}
    return ReleaseSelfTestResult(
        ok=required_states <= set(states) and boundaries and not percentages,
        states=states,
        page_boundaries_preserved=boundaries,
        user_facing_percentages=percentages,
    )
